#!/usr/bin/env python3
"""
将时序图分析结果(JSON)写入Excel表格并生成图片。
用法: python3 generate_excel_charts.py <profile_results_dir> [output_dir]
"""
import sys
import os
import json
import subprocess
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np

def run_analyze(profile_dir):
    """运行analyze脚本获取JSON结果"""
    analyze_script = str(Path(__file__).resolve().parent / "analyze_c50_profile.mjs")
    result = subprocess.run(
        ["node", analyze_script, profile_dir],
        capture_output=True, text=True, timeout=30
    )
    if result.returncode != 0:
        print(f"analyze failed: {result.stderr}", file=sys.stderr)
        sys.exit(1)
    return json.loads(result.stdout)

def style_header(ws, row, cols, fill_color="4472C4"):
    """设置表头样式"""
    fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    font = Font(bold=True, color="FFFFFF", size=11)
    align = Alignment(horizontal="center", vertical="center")
    border = Border(
        bottom=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000")
    )
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = align
        cell.border = border

def auto_width(ws):
    """自动调整列宽"""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                val = str(cell.value) if cell.value else ""
                if len(val) > max_len:
                    max_len = len(val)
            except:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 30)

def write_benchmark_sheet(wb, data):
    """写入bench总体结果"""
    ws = wb.create_sheet("Bench总体", 0)
    
    bench = data.get("benchmark", {})
    create = bench.get("create", {})
    
    headers = ["指标", "值"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    style_header(ws, 1, 2)
    
    rows = [
        ("case", data.get("case", "")),
        ("successful", bench.get("successful", "")),
        ("errors", bench.get("errors", "")),
        ("success_rate", f"{bench.get('success_rate', 0)*100:.1f}%"),
        ("throughput_qps", f"{bench.get('throughput_qps', 0):.1f}"),
        ("total_time_s", f"{bench.get('total_time_s', 0):.3f}"),
        ("", ""),
        ("create latency", ""),
        ("min", f"{create.get('min', 0):.1f}ms"),
        ("avg", f"{create.get('avg', 0):.1f}ms"),
        ("std", f"{create.get('std', 0):.1f}ms"),
        ("P50", f"{create.get('p50', 0):.1f}ms"),
        ("P90", f"{create.get('p90', 0):.1f}ms"),
        ("P95", f"{create.get('p95', 0):.1f}ms"),
        ("P99", f"{create.get('p99', 0):.1f}ms"),
        ("max", f"{create.get('max', 0):.1f}ms"),
    ]
    
    for row_idx, (label, value) in enumerate(rows, 2):
        ws.cell(row=row_idx, column=1, value=label)
        ws.cell(row=row_idx, column=2, value=value)
    
    auto_width(ws)

def write_blocks_sheet(wb, data):
    """写入分批结果"""
    ws = wb.create_sheet("分批时延")
    
    blocks = data.get("benchmark_blocks", [])
    headers = ["批次", "请求数", "avg(ms)", "min(ms)", "P50(ms)", "P95(ms)", "max(ms)"]
    
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    style_header(ws, 1, len(headers))
    
    for row_idx, block in enumerate(blocks, 2):
        ws.cell(row=row_idx, column=1, value=block.get("sequence", ""))
        ws.cell(row=row_idx, column=2, value=block.get("count", ""))
        ws.cell(row=row_idx, column=3, value=round(block.get("avg", 0), 1))
        ws.cell(row=row_idx, column=4, value=round(block.get("min", 0), 1))
        ws.cell(row=row_idx, column=5, value=round(block.get("p50", 0), 1))
        ws.cell(row=row_idx, column=6, value=round(block.get("p95", 0), 1))
        ws.cell(row=row_idx, column=7, value=round(block.get("max", 0), 1))
    
    auto_width(ws)

def write_cubelet_stages_sheet(wb, data):
    """写入cubelet各阶段"""
    ws = wb.create_sheet("Cubelet阶段")
    
    stages = data.get("cubelet_stages", {})
    headers = ["阶段", "count", "avg(ms)", "min(ms)", "P50(ms)", "P95(ms)", "max(ms)"]
    
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    style_header(ws, 1, len(headers))
    
    row_idx = 2
    for name, s in stages.items():
        if s.get("count", 0) <= 0:
            continue
        ws.cell(row=row_idx, column=1, value=name)
        ws.cell(row=row_idx, column=2, value=s.get("count", ""))
        ws.cell(row=row_idx, column=3, value=round(s.get("avg", 0), 1))
        ws.cell(row=row_idx, column=4, value=round(s.get("min", 0), 1))
        ws.cell(row=row_idx, column=5, value=round(s.get("p50", 0), 1))
        ws.cell(row=row_idx, column=6, value=round(s.get("p95", 0), 1))
        ws.cell(row=row_idx, column=7, value=round(s.get("max", 0), 1))
        row_idx += 1
    
    auto_width(ws)

def write_shim_stages_sheet(wb, data):
    """写入shim各阶段"""
    ws = wb.create_sheet("Shim阶段")
    
    stages = data.get("shim_stages", {})
    headers = ["阶段", "count", "avg(ms)", "min(ms)", "P50(ms)", "P95(ms)", "max(ms)"]
    
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    style_header(ws, 1, len(headers))
    
    row_idx = 2
    for name, s in stages.items():
        if s.get("count", 0) <= 0:
            continue
        ws.cell(row=row_idx, column=1, value=name)
        ws.cell(row=row_idx, column=2, value=s.get("count", ""))
        ws.cell(row=row_idx, column=3, value=round(s.get("avg", 0), 1))
        ws.cell(row=row_idx, column=4, value=round(s.get("min", 0), 1))
        ws.cell(row=row_idx, column=5, value=round(s.get("p50", 0), 1))
        ws.cell(row=row_idx, column=6, value=round(s.get("p95", 0), 1))
        ws.cell(row=row_idx, column=7, value=round(s.get("max", 0), 1))
        row_idx += 1
    
    auto_width(ws)

def write_detail_sheet(wb, data):
    """写入详细子阶段"""
    ws = wb.create_sheet("详细子阶段")
    
    detail = data.get("create_container_detail", {})
    timings = detail.get("direct_timings_ms", {})
    
    headers = ["子阶段", "count", "avg(ms)", "min(ms)", "P50(ms)", "P95(ms)", "max(ms)"]
    
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    style_header(ws, 1, len(headers))
    
    row_idx = 2
    for name, t in timings.items():
        if t.get("count", 0) <= 0:
            continue
        ws.cell(row=row_idx, column=1, value=name)
        ws.cell(row=row_idx, column=2, value=t.get("count", ""))
        ws.cell(row=row_idx, column=3, value=round(t.get("avg", 0), 1))
        ws.cell(row=row_idx, column=4, value=round(t.get("min", 0), 1))
        ws.cell(row=row_idx, column=5, value=round(t.get("p50", 0), 1))
        ws.cell(row=row_idx, column=6, value=round(t.get("p95", 0), 1))
        ws.cell(row=row_idx, column=7, value=round(t.get("max", 0), 1))
        row_idx += 1
    
    auto_width(ws)

def write_cubemaster_sheet(wb, data):
    """写入cubemaster各阶段"""
    ws = wb.create_sheet("CubeMaster阶段")
    
    stages = data.get("cubemaster_stages", {})
    headers = ["阶段", "count", "avg(ms)", "min(ms)", "P50(ms)", "P95(ms)", "max(ms)"]
    
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    style_header(ws, 1, len(headers))
    
    row_idx = 2
    for name, s in stages.items():
        if s.get("count", 0) <= 0:
            continue
        ws.cell(row=row_idx, column=1, value=name)
        ws.cell(row=row_idx, column=2, value=s.get("count", ""))
        ws.cell(row=row_idx, column=3, value=round(s.get("avg", 0), 1))
        ws.cell(row=row_idx, column=4, value=round(s.get("min", 0), 1))
        ws.cell(row=row_idx, column=5, value=round(s.get("p50", 0), 1))
        ws.cell(row=row_idx, column=6, value=round(s.get("p95", 0), 1))
        ws.cell(row=row_idx, column=7, value=round(s.get("max", 0), 1))
        row_idx += 1
    
    auto_width(ws)

def write_host_sheet(wb, data):
    """写入host资源"""
    ws = wb.create_sheet("Host资源")
    
    summary = data.get("host_summary", {})
    headers = ["指标", "值"]
    
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    style_header(ws, 1, 2)
    
    rows = [
        ("sample_count", summary.get("sample_count", "")),
        ("max_runnable", summary.get("max_runnable", "")),
        ("max_procs_running", summary.get("max_procs_running", "")),
        ("max_procs_blocked", summary.get("max_procs_blocked", "")),
        ("max_threads", summary.get("max_threads", "")),
        ("max_shims", summary.get("max_shims", "")),
        ("max_iowait_pct", f"{summary.get('max_iowait_pct', 0):.3f}%"),
        ("max_system_pct", f"{summary.get('max_system_pct', 0):.2f}%"),
        ("max_ctxt_per_s", f"{summary.get('max_ctxt_per_s', 0):.0f}"),
        ("max_processes_per_s", f"{summary.get('max_processes_per_s', 0):.0f}"),
    ]
    
    for row_idx, (label, value) in enumerate(rows, 2):
        ws.cell(row=row_idx, column=1, value=label)
        ws.cell(row=row_idx, column=2, value=value)
    
    auto_width(ws)

def generate_stage_chart(data, output_dir, case_name):
    """生成各阶段avg对比柱状图"""
    fig, axes = plt.subplots(2, 2, figsize=(16, 12))
    fig.suptitle(f"Timing Stages - {case_name}", fontsize=16, fontweight="bold")
    
    charts = [
        ("Cubelet", data.get("cubelet_stages", {}), axes[0, 0]),
        ("Shim", data.get("shim_stages", {}), axes[0, 1]),
        ("CubeMaster", data.get("cubemaster_stages", {}), axes[1, 0]),
    ]
    
    for title, stages, ax in charts:
        names = []
        avgs = []
        p50s = []
        p95s = []
        for name, s in stages.items():
            if s.get("count", 0) <= 0:
                continue
            names.append(name)
            avgs.append(s.get("avg", 0))
            p50s.append(s.get("p50", 0))
            p95s.append(s.get("p95", 0))
        
        if not names:
            ax.set_title(f"{title} (no data)")
            continue
        
        x = np.arange(len(names))
        width = 0.25
        
        ax.bar(x - width, avgs, width, label="avg", color="#4472C4")
        ax.bar(x, p50s, width, label="P50", color="#ED7D31")
        ax.bar(x + width, p95s, width, label="P95", color="#A5A5A5")
        
        ax.set_title(title, fontsize=13, fontweight="bold")
        ax.set_ylabel("Latency (ms)")
        ax.set_xticks(x)
        ax.set_xticklabels(names, rotation=30, ha="right", fontsize=9)
        ax.legend()
        ax.grid(axis="y", alpha=0.3)
    
    # 详细子阶段
    detail = data.get("create_container_detail", {})
    timings = detail.get("direct_timings_ms", {})
    ax = axes[1, 1]
    names = []
    avgs = []
    p50s = []
    p95s = []
    for name, t in timings.items():
        if t.get("count", 0) <= 0:
            continue
        names.append(name)
        avgs.append(t.get("avg", 0))
        p50s.append(t.get("p50", 0))
        p95s.append(t.get("p95", 0))
    
    if names:
        x = np.arange(len(names))
        width = 0.25
        ax.bar(x - width, avgs, width, label="avg", color="#4472C4")
        ax.bar(x, p50s, width, label="P50", color="#ED7D31")
        ax.bar(x + width, p95s, width, label="P95", color="#A5A5A5")
        ax.set_title("Detail Sub-stages", fontsize=13, fontweight="bold")
        ax.set_ylabel("Latency (ms)")
        ax.set_xticks(x)
        ax.set_xticklabels(names, rotation=30, ha="right", fontsize=8)
        ax.legend()
        ax.grid(axis="y", alpha=0.3)
    else:
        ax.set_title("Detail Sub-stages (no data)")
    
    plt.tight_layout()
    chart_path = os.path.join(output_dir, f"stages_chart_{case_name}.png")
    plt.savefig(chart_path, dpi=150, bbox_inches="tight")
    plt.close()
    print(f"Chart saved: {chart_path}")

def generate_blocks_chart(data, output_dir, case_name):
    """生成分批时延折线图"""
    blocks = data.get("benchmark_blocks", [])
    if not blocks:
        return
    
    fig, ax = plt.subplots(figsize=(14, 6))
    
    sequences = [b.get("sequence", "") for b in blocks]
    avgs = [b.get("avg", 0) for b in blocks]
    p50s = [b.get("p50", 0) for b in blocks]
    p95s = [b.get("p95", 0) for b in blocks]
    maxs = [b.get("max", 0) for b in blocks]
    
    x = np.arange(len(sequences))
    
    ax.plot(x, avgs, "o-", label="avg", color="#4472C4", linewidth=2)
    ax.plot(x, p50s, "s-", label="P50", color="#ED7D31", linewidth=2)
    ax.plot(x, p95s, "^-", label="P95", color="#A5A5A5", linewidth=2)
    ax.plot(x, maxs, "d-", label="max", color="#FFC000", linewidth=1.5)
    
    ax.set_title(f"Latency Timeline by Batch - {case_name}", fontsize=14, fontweight="bold")
    ax.set_xlabel("Batch (sequence)")
    ax.set_ylabel("Latency (ms)")
    ax.set_xticks(x)
    ax.set_xticklabels(sequences, rotation=45, ha="right", fontsize=9)
    ax.legend()
    ax.grid(alpha=0.3)
    
    plt.tight_layout()
    chart_path = os.path.join(output_dir, f"blocks_timeline_{case_name}.png")
    plt.savefig(chart_path, dpi=150, bbox_inches="tight")
    plt.close()
    print(f"Chart saved: {chart_path}")

def generate_host_chart(data, output_dir, case_name):
    """生成host资源时间线图"""
    intervals = data.get("host_intervals", [])
    if not intervals:
        return
    
    fig, axes = plt.subplots(2, 2, figsize=(14, 10))
    fig.suptitle(f"Host Resources - {case_name}", fontsize=14, fontweight="bold")
    
    times = [i.get("end_s", 0) for i in intervals]
    
    # runnable + shims
    ax = axes[0, 0]
    runnable = [i.get("runnable", 0) for i in intervals]
    shims = [i.get("shim_count", 0) for i in intervals]
    ax.plot(times, runnable, "o-", label="runnable", color="#4472C4")
    ax.plot(times, shims, "s-", label="shims", color="#ED7D31")
    ax.set_title("Runnable threads & Shim count")
    ax.set_xlabel("Time (s)")
    ax.set_ylabel("Count")
    ax.legend()
    ax.grid(alpha=0.3)
    
    # CPU usage
    ax = axes[0, 1]
    busy = [i.get("busy_pct", 0) for i in intervals]
    system = [i.get("system_pct", 0) for i in intervals]
    iowait = [i.get("iowait_pct", 0) * 10 for i in intervals]  # scale for visibility
    ax.plot(times, busy, "o-", label="busy%", color="#4472C4")
    ax.plot(times, system, "s-", label="system%", color="#ED7D31")
    ax.plot(times, iowait, "^-", label="iowait%x10", color="#A5A5A5")
    ax.set_title("CPU Usage")
    ax.set_xlabel("Time (s)")
    ax.set_ylabel("Percentage")
    ax.legend()
    ax.grid(alpha=0.3)
    
    # Context switches
    ax = axes[1, 0]
    ctxt = [i.get("ctxt_per_s", 0) / 1000 for i in intervals]  # K/s
    ax.plot(times, ctxt, "o-", label="ctxt K/s", color="#4472C4")
    ax.set_title("Context Switches")
    ax.set_xlabel("Time (s)")
    ax.set_ylabel("K switches/s")
    ax.legend()
    ax.grid(alpha=0.3)
    
    # Threads
    ax = axes[1, 1]
    threads = [i.get("threads", 0) / 1000 for i in intervals]  # K
    procs = [i.get("procs_running", 0) for i in intervals]
    ax.plot(times, threads, "o-", label="threads (K)", color="#4472C4")
    ax.plot(times, procs, "s-", label="procs_running", color="#ED7D31")
    ax.set_title("Threads & Running Processes")
    ax.set_xlabel("Time (s)")
    ax.set_ylabel("Count")
    ax.legend()
    ax.grid(alpha=0.3)
    
    plt.tight_layout()
    chart_path = os.path.join(output_dir, f"host_resources_{case_name}.png")
    plt.savefig(chart_path, dpi=150, bbox_inches="tight")
    plt.close()
    print(f"Chart saved: {chart_path}")

def main():
    if len(sys.argv) < 2:
        print("Usage: python3 generate_excel_charts.py <profile_results_dir> [output_dir]")
        sys.exit(1)
    
    profile_dir = sys.argv[1]
    output_dir = sys.argv[2] if len(sys.argv) > 2 else os.path.dirname(profile_dir)
    
    os.makedirs(output_dir, exist_ok=True)
    
    print(f"Analyzing: {profile_dir}")
    data = run_analyze(profile_dir)
    
    case_name = data.get("case", "unknown")
    print(f"Case: {case_name}")
    
    # 生成Excel
    wb = Workbook()
    wb.remove(wb.active)  # 删除默认sheet
    
    write_benchmark_sheet(wb, data)
    write_blocks_sheet(wb, data)
    write_cubelet_stages_sheet(wb, data)
    write_shim_stages_sheet(wb, data)
    write_detail_sheet(wb, data)
    write_cubemaster_sheet(wb, data)
    write_host_sheet(wb, data)
    
    excel_path = os.path.join(output_dir, f"timing_report_{case_name}.xlsx")
    wb.save(excel_path)
    print(f"Excel saved: {excel_path}")
    
    # 生成图片
    generate_stage_chart(data, output_dir, case_name)
    generate_blocks_chart(data, output_dir, case_name)
    generate_host_chart(data, output_dir, case_name)
    
    print(f"\nDone! Output in: {output_dir}")
    print(f"  Excel: timing_report_{case_name}.xlsx")
    print(f"  Charts: stages_chart_{case_name}.png, blocks_timeline_{case_name}.png, host_resources_{case_name}.png")

if __name__ == "__main__":
    main()
