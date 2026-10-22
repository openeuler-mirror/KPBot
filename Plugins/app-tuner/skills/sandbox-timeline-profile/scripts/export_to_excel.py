#!/usr/bin/env python3
"""Convert CubeSandbox V2 profiling results to Excel.

Usage:
    python export_to_excel.py <analysis.json> [<stage-latency.csv>] [output.xlsx]

    If stage-latency.csv is omitted, reads from the same directory as analysis.json.
"""

import csv, json, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# ── config ──────────────────────────────────────────────
KEEP_COLS = [
    "depth", "component", "stage_id", "stage_name", "parent_stage_id",
    "count", "avg_ms", "min_ms", "p50_ms", "p95_ms", "max_ms", "notes",
]

COL_WIDTHS = {
    "depth": 6, "component": 12, "stage_id": 30, "stage_name": 26,
    "parent_stage_id": 30, "count": 6, "avg_ms": 10, "min_ms": 10,
    "p50_ms": 10, "p95_ms": 10, "max_ms": 10, "notes": 60,
}

# ── args ────────────────────────────────────────────────
if len(sys.argv) < 2:
    print("Usage: python export_to_excel.py ANALYSIS_JSON [STAGE_LATENCY_CSV] [OUTPUT.xlsx]")
    sys.exit(2)

analysis_path = Path(sys.argv[1])
csv_path = Path(sys.argv[2]) if len(sys.argv) > 2 else analysis_path.parent / "stage-latency.csv"
output_path = Path(sys.argv[3]) if len(sys.argv) > 3 else analysis_path.parent / "profile-results.xlsx"

# ── load data ───────────────────────────────────────────
with open(analysis_path, encoding="utf-8") as f:
    data = json.load(f)

rows = []
if csv_path.exists():
    with open(csv_path, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
else:
    print(f"Warning: {csv_path} not found, stage-latency sheet will be empty")

# ── build workbook ──────────────────────────────────────
wb = Workbook()

# Sheet 1: Stage Latency
ws = wb.active
ws.title = "Stage Latency"
for ci, col in enumerate(KEEP_COLS, 1):
    ws.cell(row=1, column=ci, value=col)
for ri, row in enumerate(rows, 2):
    for ci, col in enumerate(KEEP_COLS, 1):
        val = row.get(col, "")
        if col == "count":
            try: val = int(float(val))
            except: pass
        elif col in ("avg_ms", "min_ms", "p50_ms", "p95_ms", "max_ms"):
            try: val = round(float(val), 2)
            except: pass
        elif col == "depth":
            try: val = int(val)
            except: pass
        ws.cell(row=ri, column=ci, value=val)
for ci, col in enumerate(KEEP_COLS, 1):
    ws.column_dimensions[get_column_letter(ci)].width = COL_WIDTHS.get(col, 12)
ws.freeze_panes = "A2"

# Sheet 2: Summary
ws2 = wb.create_sheet("Summary")
ws2.column_dimensions["A"].width = 30
ws2.column_dimensions["B"].width = 20
b, c = data["benchmark"], data["benchmark"]["create"]
items = [
    ("Case", data["case"]),
    ("Correlated Instances", data["correlated_instances"]),
    ("", ""),
    ("Errors", b["errors"]),
    ("Success Rate", f"{b['success_rate'] * 100}%"),
    ("Throughput (QPS)", round(b["throughput_qps"], 1)),
    ("Total Time (s)", round(b["total_time_s"], 2)),
    ("Create Avg (ms)", round(c["avg"], 2)),
    ("Create P50 (ms)", round(c["p50"], 2)),
    ("Create P95 (ms)", round(c["p95"], 2)),
    ("Create P99 (ms)", round(c["p99"], 2)),
    ("Create Max (ms)", round(c["max"], 2)),
    ("", ""),
    ("Max Runnable", data["host_summary"]["max_runnable"]),
    ("Max Shims", data["host_summary"]["max_shims"]),
    ("Max IOWait %", round(data["host_summary"]["max_iowait_pct"], 3)),
    ("Max System %", round(data["host_summary"]["max_system_pct"], 2)),
    ("", ""),
    ("Guest path counts", ""),
]
cd = data.get("create_container_detail", {})
for path, cnt in cd.get("path_counts", {}).items():
    items.append((f"  {path}", cnt))
for metric, value in items:
    ws2.append([metric, value])

wb.save(output_path)
print(f"Saved: {output_path}  ({len(rows)} stage rows)")
